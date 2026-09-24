import streamlit as st
import pandas as pd
import numpy as np
import plotly.express as px
import os
import io
import datetime
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter

# 1. Configuración del Dashboard
st.set_page_config(page_title="Visualizador de Curva Cupón Cero (VTD)", layout="wide")
st.title("Dashboard de Vector de Tasa de Descuento Ajustado")

archivos_fuente = {
    "Soles (CSS)": "curva_historica_soles.xlsx",
    "Soles VAC (CSS VAC)": "curva_historica_soles_vac.xlsx",
    "Dólares (CDG)": "curva_historica_dolares.xlsx"
}

# Valores VOLA predeterminados por moneda
VOLA_VALORES = {
    "Soles (CSS)": 0.001502,      # 0.1502%
    "Soles VAC (CSS VAC)": 0.002471,  # 0.2471%
    "Dólares (CDG)": 0.002175     # 0.2175%
}

colores_graficos = ["#06369D", "#ED8B00", "#1FC3B3", "#97D700", "#981D97", "#BEBAB9"]

# Inicialización de variables en memoria
if 'tablas_mostrar' not in st.session_state:
    st.session_state.tablas_mostrar = {}
if 'datos_grafico' not in st.session_state:
    st.session_state.datos_grafico = []
if 'comp_resultados' not in st.session_state:
    st.session_state.comp_resultados = {}
if 'comp_df_list' not in st.session_state:
    st.session_state.comp_df_list = []

# --- 2A. CARGA EN MEMORIA ---
@st.cache_data
def cargar_datos_en_memoria(ruta_archivo):
    """Lee el archivo crudo y lo deja pivoteado y listo en RAM."""
    try:
        if ruta_archivo.endswith('.parquet'):
            df_pivot = pd.read_parquet(ruta_archivo)
        else:
            df_raw = pd.read_excel(ruta_archivo, header=1)
            df_raw['Fecha de Proceso'] = pd.to_datetime(df_raw['Fecha de Proceso'], format='%d/%m/%Y')
            df_pivot = df_raw.pivot(index='Fecha de Proceso', columns='Plazo (DIAS)', values='Tasas (%)')

        plazos_completos = list(range(0, 14401, 90))
        df_pivot = df_pivot.reindex(columns=plazos_completos).ffill(axis=1)
        return df_pivot
    except Exception:
        return None

bases_de_datos = {}
for moneda, ruta in archivos_fuente.items():
    if os.path.exists(ruta):
        bases_de_datos[moneda] = cargar_datos_en_memoria(ruta)

# --- 2B. MOTOR DE CÁLCULO SPOT & FORWARD ---
def calcular_curva_cupon_cero(df_pivot, fecha_inicio, fecha_fin):
    """Calcula la curva spot interpolada a 480 meses."""
    f_inicio = pd.to_datetime(fecha_inicio)
    f_fin = pd.to_datetime(fecha_fin)
    df_filtrado = df_pivot.loc[(df_pivot.index >= f_inicio) & (df_pivot.index <= f_fin)]

    if df_filtrado.empty:
        return None

    tasas_promedio = df_filtrado.mean().values
    meses = np.arange(481)
    df_curva = pd.DataFrame({'Mes': meses})
    df_curva['Año'] = df_curva['Mes'] // 12

    block_idx = df_curva['Mes'] // 3
    idx_inicio = np.clip(block_idx, 0, len(tasas_promedio) - 1)
    idx_fin = np.clip(block_idx + 1, 0, len(tasas_promedio) - 1)

    df_curva['Inicio'] = tasas_promedio[idx_inicio]
    df_curva['Fin'] = tasas_promedio[idx_fin]
    df_curva['Flag Inicio'] = block_idx * 3
    df_curva['Flag Fin'] = (block_idx + 1) * 3

    df_curva['Tasa anual (%)'] = (
        df_curva['Fin'] - ((df_curva['Fin'] - df_curva['Inicio']) * (df_curva['Flag Fin'] - df_curva['Mes']) / 3)
    ) / 100

    df_curva['Tasa mensual (%)'] = (1 + df_curva['Tasa anual (%)']) ** (1/12) - 1
    return df_curva[['Año', 'Tasa anual (%)', 'Mes', 'Tasa mensual (%)']]

def calcular_tasas_forward(df_curva, col_spot='Tasa anual + VOLA (%)'):
    """
    Calcula la curva forward a partir de la tasa Spot.
    """
    df = df_curva.copy()
    s = df[col_spot].values
    t = df['Mes'].values
    
    forward = np.zeros_like(s)
    forward[0] = s[0]
    
    factor_acum = (1.0 + s) ** t
    forward[1:] = (factor_acum[1:] / factor_acum[:-1]) - 1.0
    
    df['Tasa Forward (%)'] = forward
    return df

# --- 2C. GENERADORES DE EXCEL ---
def generar_excel_vtd_vola(dict_tasas, fecha_inicio, fecha_fin, valores_vola):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "VTD + VOLA"
    ws.views.sheetView[0].showGridLines = True

    font_title = Font(name="Times New Roman", size=20, bold=True)
    font_subtitle = Font(name="Times New Roman", size=14, bold=True)
    font_header_group = Font(name="Times New Roman", size=11, bold=True)
    font_header = Font(name="Times New Roman", size=10, bold=True)
    font_data = Font(name="Arial Narrow", size=10)
    font_vola_val = Font(name="Arial Narrow", size=10)

    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_right = Alignment(horizontal="right", vertical="center")
    border_thin_bottom = Border(bottom=Side(style='thin', color='000000'))
    border_last_row = Border(bottom=Side(style='medium', color='000000'))

    ws.merge_cells("B2:P2")
    ws["B2"] = "Vector de tasa de descuento (VTD) y ajuste por volatilidad (VOLA)"
    ws["B2"].font = font_title
    ws["B2"].alignment = align_center

    ws.merge_cells("B3:P3")
    ws["B3"] = "(tasas anuales)"
    ws["B3"].font = font_subtitle
    ws["B3"].alignment = align_center

    ws.row_dimensions[2].height = 25.8
    ws.row_dimensions[6].height = 34

    meses_es = {1: 'Enero', 2: 'Febrero', 3: 'Marzo', 4: 'Abril', 5: 'Mayo', 6: 'Junio', 7: 'Julio', 8: 'Agosto', 9: 'Septiembre', 10: 'Octubre', 11: 'Noviembre', 12: 'Diciembre'}
    f_ini_str = f"{meses_es[fecha_inicio.month]}-{fecha_inicio.year}"
    f_fin_str = f"{meses_es[fecha_fin.month]}-{fecha_fin.year}"
    texto_periodo = f"VTD aplicable a {f_fin_str}\n(Promedio {f_ini_str} a {f_fin_str})"

    ws.merge_cells("B6:F6")
    ws["B6"] = texto_periodo
    ws["B6"].font = font_header_group
    ws["B6"].alignment = align_center

    for col in range(2, 7):
        ws.cell(row=6, column=col).border = border_thin_bottom

    headers_t1 = ["Mes", "Año", "Soles", "Soles VAC", "Dólares"]
    cols_t1 = ["B", "C", "D", "E", "F"]
    for col_letter, header in zip(cols_t1, headers_t1):
        cell = ws[f"{col_letter}8"]
        cell.value = header
        cell.font = font_header
        cell.alignment = align_center
        cell.border = border_thin_bottom

    ws.merge_cells("H6:L6")
    ws["H6"] = "VTD incluyendo VOLA"
    ws["H6"].font = font_header_group
    ws["H6"].alignment = align_center

    for col in range(8, 13):
        ws.cell(row=6, column=col).border = border_thin_bottom

    headers_t2 = ["Mes", "Año", "Soles", "Soles VAC", "Dólares"]
    cols_t2 = ["H", "I", "J", "K", "L"]
    for col_letter, header in zip(cols_t2, headers_t2):
        cell = ws[f"{col_letter}8"]
        cell.value = header
        cell.font = font_header
        cell.alignment = align_center
        cell.border = border_thin_bottom

    ws.merge_cells("N6:P6")
    ws["N6"] = "VOLA"
    ws["N6"].font = font_header_group
    ws["N6"].alignment = align_center

    for col in range(14, 17):
        ws.cell(row=6, column=col).border = border_thin_bottom

    headers_t3 = ["Soles", "Soles VAC", "Dólares"]
    cols_t3 = ["N", "O", "P"]
    for col_letter, header in zip(cols_t3, headers_t3):
        cell = ws[f"{col_letter}8"]
        cell.value = header
        cell.font = font_header
        cell.alignment = align_center
        cell.border = border_thin_bottom

    vola_soles = valores_vola.get("Soles (CSS)", 0.001502)
    vola_vac = valores_vola.get("Soles VAC (CSS VAC)", 0.002471)
    vola_usd = valores_vola.get("Dólares (CDG)", 0.002175)

    ws["N10"] = vola_soles
    ws["O10"] = vola_vac
    ws["P10"] = vola_usd

    for col_letter in cols_t3:
        cell = ws[f"{col_letter}10"]
        cell.font = font_vola_val
        cell.alignment = align_right
        cell.number_format = '0.0000%'
        cell.border = border_thin_bottom

    df_soles = dict_tasas.get("Soles (CSS)")
    df_vac = dict_tasas.get("Soles VAC (CSS VAC)")
    df_usd = dict_tasas.get("Dólares (CDG)")

    row_start = 10
    total_meses_excel = 1321
    val_soles_curr, val_vac_curr, val_usd_curr = 0, 0, 0

    for i in range(total_meses_excel):
        curr_row = row_start + i
        ws[f"B{curr_row}"] = i
        ws[f"C{curr_row}"] = i // 12
        ws[f"H{curr_row}"] = i
        ws[f"I{curr_row}"] = i // 12

        for col_l in ["B", "C", "H", "I"]:
            ws[f"{col_l}{curr_row}"].font = font_data
            ws[f"{col_l}{curr_row}"].alignment = align_center

        if i <= 480:
            val_soles_curr = df_soles['Tasa anual (%)'].iloc[i] if df_soles is not None else 0
            val_vac_curr = df_vac['Tasa anual (%)'].iloc[i] if df_vac is not None else 0
            val_usd_curr = df_usd['Tasa anual (%)'].iloc[i] if df_usd is not None else 0

        ws[f"D{curr_row}"] = val_soles_curr
        ws[f"E{curr_row}"] = val_vac_curr
        ws[f"F{curr_row}"] = val_usd_curr

        for col_l in ["D", "E", "F"]:
            cell = ws[f"{col_l}{curr_row}"]
            cell.font = font_data
            cell.alignment = align_right
            cell.number_format = '0.000%'

        ws[f"J{curr_row}"] = val_soles_curr + vola_soles
        ws[f"K{curr_row}"] = val_vac_curr + vola_vac
        ws[f"L{curr_row}"] = val_usd_curr + vola_usd

        for col_l in ["J", "K", "L"]:
            cell = ws[f"{col_l}{curr_row}"]
            cell.font = font_data
            cell.alignment = align_right
            cell.number_format = '0.000%'

        if i == total_meses_excel - 1:
            for col_l in ["B", "C", "D", "E", "F", "H", "I", "J", "K", "L"]:
                ws[f"{col_l}{curr_row}"].border = border_last_row

    column_widths = {
        'A': 3, 'B': 8, 'C': 8, 'D': 14, 'E': 14, 'F': 14, 'G': 4,
        'H': 8, 'I': 8, 'J': 14, 'K': 14, 'L': 14, 'M': 4, 'N': 14, 'O': 14, 'P': 14
    }
    for col, width in column_widths.items():
        ws.column_dimensions[col].width = width

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()

def generar_excel_tasas_forward(dict_tasas, fecha_inicio, fecha_fin):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "VTD + VOLA y Forward"
    ws.views.sheetView[0].showGridLines = True

    font_title = Font(name="Times New Roman", size=20, bold=True)
    font_subtitle = Font(name="Times New Roman", size=14, bold=True)
    font_header_group = Font(name="Times New Roman", size=11, bold=True)
    font_header = Font(name="Times New Roman", size=10, bold=True)
    font_data = Font(name="Arial Narrow", size=10)

    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_right = Alignment(horizontal="right", vertical="center")

    border_thin_bottom = Border(bottom=Side(style='thin', color='000000'))
    border_last_row = Border(bottom=Side(style='medium', color='000000'))

    ws.merge_cells("B2:L2")
    ws["B2"] = "Vector de tasa de descuento ajustado (VTDA) y tasas forward"
    ws["B2"].font = font_title
    ws["B2"].alignment = align_center

    ws.merge_cells("B3:L3")
    ws["B3"] = "(tasas anuales)"
    ws["B3"].font = font_subtitle
    ws["B3"].alignment = align_center

    ws.row_dimensions[2].height = 25.8
    ws.row_dimensions[6].height = 34

    meses_es = {1: 'Enero', 2: 'Febrero', 3: 'Marzo', 4: 'Abril', 5: 'Mayo', 6: 'Junio', 7: 'Julio', 8: 'Agosto', 9: 'Septiembre', 10: 'Octubre', 11: 'Noviembre', 12: 'Diciembre'}
    f_ini_str = f"{meses_es[fecha_inicio.month]}-{fecha_inicio.year}"
    f_fin_str = f"{meses_es[fecha_fin.month]}-{fecha_fin.year}"
    texto_periodo = f"VTDA aplicable a {f_fin_str}\n(Promedio {f_ini_str} a {f_fin_str})"

    ws.merge_cells("B6:F6")
    ws["B6"] = texto_periodo
    ws["B6"].font = font_header_group
    ws["B6"].alignment = align_center

    for col in range(2, 7):
        ws.cell(row=6, column=col).border = border_thin_bottom

    headers_t1 = ["Mes", "Año", "Soles", "Soles VAC", "Dólares"]
    cols_t1 = ["B", "C", "D", "E", "F"]
    for col_letter, header in zip(cols_t1, headers_t1):
        cell = ws[f"{col_letter}8"]
        cell.value = header
        cell.font = font_header
        cell.alignment = align_center
        cell.border = border_thin_bottom

    ws.merge_cells("H6:L6")
    ws["H6"] = "Tasas Forward (a partir de VTDA)"
    ws["H6"].font = font_header_group
    ws["H6"].alignment = align_center

    for col in range(8, 13):
        ws.cell(row=6, column=col).border = border_thin_bottom

    headers_t2 = ["Mes", "Año", "Soles", "Soles VAC", "Dólares"]
    cols_t2 = ["H", "I", "J", "K", "L"]
    for col_letter, header in zip(cols_t2, headers_t2):
        cell = ws[f"{col_letter}8"]
        cell.value = header
        cell.font = font_header
        cell.alignment = align_center
        cell.border = border_thin_bottom

    df_soles = dict_tasas.get("Soles (CSS)")
    df_vac = dict_tasas.get("Soles VAC (CSS VAC)")
    df_usd = dict_tasas.get("Dólares (CDG)")

    row_start = 10
    total_meses_excel = 1332  

    meses_indices = np.arange(total_meses_excel)
    
    def expandir_spot(df):
        if df is None:
            return np.zeros(total_meses_excel)
        spot_base = df['Tasa anual + VOLA (%)'].values 
        cola = np.full(total_meses_excel - len(spot_base), spot_base[-1])
        return np.concatenate([spot_base, cola])

    spot_soles = expandir_spot(df_soles)
    spot_vac = expandir_spot(df_vac)
    spot_usd = expandir_spot(df_usd)

    def calcular_fwd_array(s):
        fwd = np.zeros_like(s)
        fwd[0] = s[0]
        factor = (1.0 + s) ** meses_indices
        fwd[1:] = (factor[1:] / factor[:-1]) - 1.0
        return fwd

    fwd_soles = calcular_fwd_array(spot_soles)
    fwd_vac = calcular_fwd_array(spot_vac)
    fwd_usd = calcular_fwd_array(spot_usd)

    for i in range(total_meses_excel):
        curr_row = row_start + i

        ws[f"B{curr_row}"] = i
        ws[f"C{curr_row}"] = i // 12
        ws[f"H{curr_row}"] = i
        ws[f"I{curr_row}"] = i // 12

        for col_l in ["B", "C", "H", "I"]:
            ws[f"{col_l}{curr_row}"].font = font_data
            ws[f"{col_l}{curr_row}"].alignment = align_center

        ws[f"D{curr_row}"] = spot_soles[i]
        ws[f"E{curr_row}"] = spot_vac[i]
        ws[f"F{curr_row}"] = spot_usd[i]

        for col_letter in ["D", "E", "F"]:
            cell = ws[f"{col_letter}{curr_row}"]
            cell.font = font_data
            cell.alignment = align_right
            cell.number_format = '0.000%'

        ws[f"J{curr_row}"] = fwd_soles[i]
        ws[f"K{curr_row}"] = fwd_vac[i]
        ws[f"L{curr_row}"] = fwd_usd[i]

        for col_letter in ["J", "K", "L"]:
            cell = ws[f"{col_letter}{curr_row}"]
            cell.font = font_data
            cell.alignment = align_right
            cell.number_format = '0.000%'

        if i == total_meses_excel - 1:
            for col_letter in ["B", "C", "D", "E", "F", "H", "I", "J", "K", "L"]:
                ws[f"{col_letter}{curr_row}"].border = border_last_row

    column_widths = {
        'A': 3, 'B': 8, 'C': 8, 'D': 14, 'E': 14, 'F': 14, 'G': 4,
        'H': 8, 'I': 8, 'J': 14, 'K': 14, 'L': 14
    }
    for col, width in column_widths.items():
        ws.column_dimensions[col].width = width

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()

def generar_excel_prophet(dict_tasas, fecha_evaluacion, tipo_producto, mes_inicio):
    """Genera el Excel específico de Prophet emulando la macro VBA"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "VTD"
    ws.views.sheetView[0].showGridLines = False

    # 1. Expandir las tasas spot y calcular el forward
    total_meses = 1332
    meses_indices = np.arange(total_meses)
    
    def get_fwd_array(moneda_key):
        df = dict_tasas.get(moneda_key)
        if df is None:
            return np.zeros(total_meses)
        
        spot_base = df['Tasa anual + VOLA (%)'].values
        cola = np.full(total_meses - len(spot_base), spot_base[-1])
        spot_ext = np.concatenate([spot_base, cola])
        
        fwd = np.zeros_like(spot_ext)
        fwd[0] = spot_ext[0]
        factor = (1.0 + spot_ext) ** meses_indices
        fwd[1:] = (factor[1:] / factor[:-1]) - 1.0
        return fwd * 100

    fwd_soles = get_fwd_array("Soles (CSS)")
    fwd_vac = get_fwd_array("Soles VAC (CSS VAC)")
    fwd_usd = get_fwd_array("Dólares (CDG)")

    # 2. Configuración según producto
    max_cols_map = {
        "Vida Tradicional": 1224,
        "Ahorro-Inversión": 1218,
        "SPP y SCTR": 1199,
        "Renta Particular": 1199
    }
    max_cols = max_cols_map.get(tipo_producto, 1224)
    offset = 0 if mes_inicio == "Mes 0" else 1

    # Cortar los arrays según la cantidad de meses y offset requeridos
    row_vac = fwd_vac[offset : offset + max_cols].tolist()
    row_dol_nom = fwd_usd[offset : offset + max_cols].tolist()
    row_sol_aj = fwd_soles[offset : offset + max_cols].tolist()
    row_dol_aj = fwd_usd[offset : offset + max_cols].tolist()
    row_sol_nom = fwd_soles[offset : offset + max_cols].tolist()

    # 3. Generar encabezados de fecha (YYYYMM) iterativamente 
    fechas_ym = []
    dt_eval = pd.to_datetime(fecha_evaluacion)
    for i in range(max_cols):
        dt_next = dt_eval + pd.DateOffset(months=i+1)
        fechas_ym.append(int(dt_next.strftime("%Y%m")))

    # 4. Construcción gráfica en el Excel
    fill_blue = PatternFill(start_color="002060", end_color="002060", fill_type="solid")
    font_white_bold = Font(color="FFFFFF", bold=True)
    font_bold = Font(bold=True)
    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    border_thin = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))

    # Banner superior
    ws.merge_cells("A2:F8")
    ws["A2"] = "Obtención de Tasas Forward para Prophet"
    ws["A2"].fill = fill_blue
    ws["A2"].font = Font(color="FFFFFF", size=18, bold=True)
    ws["A2"].alignment = align_center

    # Cuadro de Inputs
    ws["D10"] = "Fecha de evaluación"
    ws["D11"] = "Tipo de producto"
    ws["D12"] = "Mes de Inicio"
    
    ws["E10"] = fecha_evaluacion.strftime("%d/%m/%Y")
    ws["E11"] = tipo_producto
    ws["E12"] = mes_inicio

    for r in range(10, 13):
        ws[f"D{r}"].fill = fill_blue
        ws[f"D{r}"].font = font_white_bold
        ws[f"D{r}"].alignment = align_left
        ws[f"D{r}"].border = border_thin
        ws[f"E{r}"].border = border_thin
        ws[f"E{r}"].alignment = align_left

    # Encabezados de Tabla (Moneda | Tipo de Moneda | YYYYMM | ...)
    ws["C14"] = "Moneda"
    ws["D14"] = "Tipo de Moneda"
    ws["C14"].fill = fill_blue
    ws["C14"].font = font_white_bold
    ws["C14"].border = border_thin
    ws["C14"].alignment = align_center

    ws["D14"].fill = fill_blue
    ws["D14"].font = font_white_bold
    ws["D14"].border = border_thin
    ws["D14"].alignment = align_center

    for col_idx, ym in enumerate(fechas_ym):
        c = ws.cell(row=14, column=5 + col_idx) # Inicia en la columna E (5)
        c.value = ym
        c.fill = fill_blue
        c.font = font_white_bold
        c.alignment = align_center
        c.border = border_thin

    # Inserción de Datos (Filas 15 a 19)
    filas_datos = [
        (1, "Soles VAC", row_vac),
        (2, "Dólares Nominales", row_dol_nom),
        (3, "Soles Ajustados", row_sol_aj),
        (4, "Dólares Ajustados", row_dol_aj),
        (5, "Soles Nominales", row_sol_nom),
    ]

    # Excepción explícita: Si es SPP y SCTR, borra la fila de Soles Nominales (Row 19)
    if tipo_producto == "SPP y SCTR":
        filas_datos[4] = ("", "", [""] * max_cols)

    for r_idx, (mon_id, mon_nombre, datos) in enumerate(filas_datos):
        row_num = 15 + r_idx
        
        c_id = ws.cell(row=row_num, column=3, value=mon_id)
        c_id.font = font_bold
        c_id.alignment = align_center
        if mon_id != "":
            c_id.border = border_thin

        c_nom = ws.cell(row=row_num, column=4, value=mon_nombre)
        c_nom.font = font_bold
        c_nom.alignment = align_left
        if mon_nombre != "":
            c_nom.border = border_thin

        for col_idx, val in enumerate(datos):
            c_val = ws.cell(row=row_num, column=5 + col_idx, value=val)
            if val != "":
                c_val.border = border_thin

    # Ajuste manual de anchos de columna
    ws.column_dimensions['A'].width = 3
    ws.column_dimensions['B'].width = 12
    ws.column_dimensions['C'].width = 15
    ws.column_dimensions['D'].width = 25

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()

# ==========================================
# CREACIÓN DE PESTAÑAS (TABS)
# ==========================================
tab_general, tab_forward, tab_comparativa, tab_vola_hist, tab_vpn = st.tabs([
    "📈 Análisis General (Spot)", 
    "🔮 Tasas Forward", 
    "⚖️ Comparativa de VTDA (Spot)",
    "🗃️ Histórico VOLA",
    "💰 VPN (Flujos)"
])

# ==========================================
# PESTAÑA 1: ANÁLISIS GENERAL (SPOT)
# ==========================================
with tab_general:
    if os.path.exists("logo.jpg"):
        st.sidebar.image("logo.jpg", use_container_width=True)
    
    st.sidebar.header("Parámetros Análisis General")

    fecha_inicio = st.sidebar.date_input(
        "Fecha Inicio",
        value=datetime.date(2025, 1, 1),
        min_value=datetime.date(2015, 1, 1),
        max_value=datetime.date(2030, 12, 31),
        key="spot_f_ini"
    )

    fecha_fin = st.sidebar.date_input(
        "Fecha Fin",
        value=datetime.date(2025, 12, 31),
        min_value=datetime.date(2015, 1, 1),
        max_value=datetime.date(2030, 12, 31),
        key="spot_f_fin"
    )

    monedas_seleccionadas = st.sidebar.multiselect(
        "Seleccione las curvas a graficar:",
        options=list(archivos_fuente.keys()),
        default=["Soles (CSS)"],
        key="spot_monedas"
    )

    st.sidebar.markdown("---")
    st.sidebar.subheader("⚙️ Ajuste por Volatilidad (VOLA)")
    vola_soles_in = st.sidebar.number_input("VOLA Soles (%)", value=0.1502, format="%.4f") / 100
    vola_vac_in = st.sidebar.number_input("VOLA Soles VAC (%)", value=0.2471, format="%.4f") / 100
    vola_usd_in = st.sidebar.number_input("VOLA Dólares (%)", value=0.2175, format="%.4f") / 100

    volas_usuario = {
        "Soles (CSS)": vola_soles_in,
        "Soles VAC (CSS VAC)": vola_vac_in,
        "Dólares (CDG)": vola_usd_in
    }

    st.sidebar.markdown("---")
    st.sidebar.subheader("📊 Disponibilidad de Datos")

    for moneda, df_pivot in bases_de_datos.items():
        if df_pivot is not None:
            f_min, f_max = df_pivot.index.min(), df_pivot.index.max()
            st.sidebar.info(f"**{moneda}**\n\nDesde: {f_min.strftime('%d/%m/%Y')}\nHasta: {f_max.strftime('%d/%m/%Y')}")
        else:
            st.sidebar.error(f"**{moneda}:** Error al cargar datos")

    st.sidebar.markdown("---")
    st.sidebar.markdown("**¿Faltan datos recientes?**")
    st.sidebar.link_button(
        label="🌐 Descargar nuevas tasas aquí",
        url="https://www.sbs.gob.pe/app/pp/n_CurvaSoberana/CurvaSoberana/ConsultaHistorica",
        use_container_width=True
    )

    if st.button("Generar curvas de la VTD", key="btn_gen_spot"):
        st.session_state.tablas_mostrar = {}
        st.session_state.datos_grafico = []

        if not monedas_seleccionadas:
            st.warning("Por favor, seleccione al menos una moneda.")
        else:
            for moneda in monedas_seleccionadas:
                if moneda in bases_de_datos and bases_de_datos[moneda] is not None:
                    df_resultado = calcular_curva_cupon_cero(bases_de_datos[moneda], fecha_inicio, fecha_fin)

                    if df_resultado is not None:
                        vola_aplicable = volas_usuario.get(moneda, 0)
                        df_resultado['Tasa anual + VOLA (%)'] = df_resultado['Tasa anual (%)'] + vola_aplicable
                        st.session_state.tablas_mostrar[moneda] = df_resultado

                        df_plot = df_resultado[['Año', 'Mes', 'Tasa anual + VOLA (%)']].copy()
                        df_plot['Moneda'] = moneda
                        st.session_state.datos_grafico.append(df_plot)
                else:
                    st.warning(f"No hay datos en {moneda} para el rango de fechas seleccionado.")

    if st.session_state.datos_grafico:
        df_consolidado = pd.concat(st.session_state.datos_grafico)
        df_consolidado['Tasa Graficar (%)'] = df_consolidado['Tasa anual + VOLA (%)'] * 100

        vista_eje_x = st.radio(
            "Seleccione la escala del gráfico:",
            ["Mensual", "Anual"],
            horizontal=True,
            key="vista_eje_gen"
        )

        if vista_eje_x == "Mensual":
            df_plot_final = df_consolidado[df_consolidado['Mes'] <= 240]
            col_eje_x = "Mes"
            etiqueta_x = "Plazo (Meses)"
        else:
            df_plot_final = df_consolidado[(df_consolidado['Mes'] % 12 == 0) & (df_consolidado['Año'] <= 20)]
            col_eje_x = "Año"
            etiqueta_x = "Plazo (Años)"

        fig = px.line(
            df_plot_final,
            x=col_eje_x,
            y="Tasa Graficar (%)",
            color="Moneda",
            title="Curva VTDA (Tasa anual + VOLA)",
            labels={"Tasa Graficar (%)": "Tasa Anual + VOLA (%)", col_eje_x: etiqueta_x},
            template="plotly_white",
            color_discrete_sequence=colores_graficos,
            markers=True if col_eje_x == "Año" else False
        )

        fig.update_layout(
            yaxis_tickformat='.2f',
            legend=dict(orientation="h", yanchor="bottom", y=1.02, xanchor="right", x=1),
            margin=dict(l=10, r=10, t=50, b=10)
        )

        if col_eje_x == "Año":
            fig.update_xaxes(dtick=1, range=[0, 20])
        else:
            fig.update_xaxes(dtick=12, range=[0, 240])

        st.plotly_chart(fig, use_container_width=True)
        
        st.subheader("📥 Descarga de reportes Excel")
        st.markdown("Se obtiene información de los 3 tipos de monedas.")

        str_inicio = fecha_inicio.strftime('%d-%m-%Y')
        str_fin = fecha_fin.strftime('%d-%m-%Y')

        dict_3_monedas = {}
        for m_key in archivos_fuente.keys():
            if m_key in bases_de_datos and bases_de_datos[m_key] is not None:
                dict_3_monedas[m_key] = calcular_curva_cupon_cero(bases_de_datos[m_key], fecha_inicio, fecha_fin)

        col_down1, col_down2 = st.columns(2)

        with col_down1:
            nombre_excel_orig = f"Reporte_VTD_{str_inicio}_al_{str_fin}.xlsx"
            buffer_orig = io.BytesIO()
            with pd.ExcelWriter(buffer_orig, engine='openpyxl') as writer:
                df_resumen = pd.DataFrame({
                    'Parámetro': ['Fecha Inicio de Consulta', 'Fecha Fin de Consulta', 'Curvas Generadas'],
                    'Valor': [str_inicio, str_fin, ", ".join(st.session_state.tablas_mostrar.keys())]
                })
                df_resumen.to_excel(writer, sheet_name='Resumen_Filtros', index=False)
                for moneda, df in st.session_state.tablas_mostrar.items():
                    df.to_excel(writer, sheet_name=moneda[:31], index=False)

            st.download_button(
                label="📄 Descargar tasas mensuales (según filtro)",
                data=buffer_orig.getvalue(),
                file_name=nombre_excel_orig,
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                key="btn_descarga_individual"
            )

        with col_down2:
            excel_vtd_vola_bytes = generar_excel_vtd_vola(dict_3_monedas, fecha_inicio, fecha_fin, volas_usuario)
            nombre_vtd_vola = f"VTD_y_VOLA_Consolidado_{str_inicio}_al_{str_fin}.xlsx"

            st.download_button(
                label="📊 Descargar formato oficial SBS (VTD + VOLA)",
                data=excel_vtd_vola_bytes,
                file_name=nombre_vtd_vola,
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                key="btn_descarga_vtd_vola"
            )

        st.markdown("---")
        cols = st.columns(len(st.session_state.tablas_mostrar))
        for i, (moneda, df) in enumerate(st.session_state.tablas_mostrar.items()):
            with cols[i]:
                st.markdown(f"**{moneda}**")
                df_display = df[['Año', 'Mes', 'Tasa anual + VOLA (%)']]
                st.dataframe(
                    df_display.style.format({'Tasa anual + VOLA (%)': '{:.4%}'}),
                    hide_index=True,
                    height=400
                )

# ==========================================
# PESTAÑA 2: TASAS FORWARD (CON VOLA) + PROPHET
# ==========================================
with tab_forward:
    st.subheader("Vector de Tasas Forward")
    st.caption("Cálculo sobre la curva VTD Spot (con VOLA): $Forward_t = \\frac{(1 + S_t)^t}{(1 + S_{t-1})^{t-1}} - 1$")

    if not st.session_state.tablas_mostrar:
        st.info("ℹ️ Por favor, genera primero las curvas en la pestaña **'📈 Análisis General'** para habilitar los cálculos forward.")
    else:
        dict_forward = {}
        forward_plot_list = []

        for moneda, df_spot in st.session_state.tablas_mostrar.items():
            df_fwd = calcular_tasas_forward(df_spot, col_spot='Tasa anual + VOLA (%)')
            dict_forward[moneda] = df_fwd

            df_p = df_fwd[['Año', 'Mes', 'Tasa Forward (%)']].copy()
            df_p['Moneda'] = moneda
            forward_plot_list.append(df_p)

        df_fwd_consolidado = pd.concat(forward_plot_list)
        df_fwd_consolidado['Tasa Graficar (%)'] = df_fwd_consolidado['Tasa Forward (%)'] * 100

        vista_eje_fwd = st.radio(
            "Seleccione la escala del gráfico:",
            ["Mensual", "Anual"],
            horizontal=True,
            key="vista_eje_fwd"
        )

        if vista_eje_fwd == "Mensual":
            df_fwd_plot_final = df_fwd_consolidado[df_fwd_consolidado['Mes'] <= 240]
            col_eje_fwd = "Mes"
            etiqueta_x_fwd = "Plazo (Meses)"
        else:
            df_fwd_plot_final = df_fwd_consolidado[(df_fwd_consolidado['Mes'] % 12 == 0) & (df_fwd_consolidado['Año'] <= 20)]
            col_eje_fwd = "Año"
            etiqueta_x_fwd = "Plazo (Años)"

        mis_colores = ["#06369D", "#ED8B00", "#1FC3B3", "#97D700", "#981D97", "#BEBAB9"]

        fig_fwd = px.line(
            df_fwd_plot_final,
            x=col_eje_fwd,
            y="Tasa Graficar (%)",
            color="Moneda",
            title="Curva de Tasas Forward (a partir de VTD Spot ajustada con VOLA)",
            labels={"Tasa Graficar (%)": "Tasa Forward (%)", col_eje_fwd: etiqueta_x_fwd},
            template="plotly_white",
            color_discrete_sequence=mis_colores,
            markers=True if col_eje_fwd == "Año" else False
        )

        fig_fwd.update_layout(
            yaxis_tickformat='.2f',
            legend=dict(orientation="h", yanchor="bottom", y=1.02, xanchor="right", x=1),
            margin=dict(l=10, r=10, t=50, b=10)
        )

        if col_eje_fwd == "Año":
            fig_fwd.update_xaxes(dtick=1, range=[0, 20])
        else:
            fig_fwd.update_xaxes(dtick=12, range=[0, 240])

        st.plotly_chart(fig_fwd, use_container_width=True)

        st.subheader("📥 Descarga de Tasas Forward")
        st.markdown("Se obtiene información de los 3 tipos de monedas.")
        str_inicio = fecha_inicio.strftime('%d-%m-%Y')
        str_fin = fecha_fin.strftime('%d-%m-%Y')

        dict_3_monedas_forward = {}
        for m_key in archivos_fuente.keys():
            if m_key in bases_de_datos and bases_de_datos[m_key] is not None:
                df_temp = calcular_curva_cupon_cero(bases_de_datos[m_key], fecha_inicio, fecha_fin)
                if df_temp is not None:
                    vola_aplicable = volas_usuario.get(m_key, 0)
                    df_temp['Tasa anual + VOLA (%)'] = df_temp['Tasa anual (%)'] + vola_aplicable
                    dict_3_monedas_forward[m_key] = df_temp

        excel_fwd_bytes = generar_excel_tasas_forward(dict_3_monedas_forward, fecha_inicio, fecha_fin)
        nombre_fwd_excel = f"VTD_VOLA_y_Tasas_Forward_Consolidado_{str_inicio}_al_{str_fin}.xlsx"

        st.download_button(
            label="📊 Descargar reporte (VTDA & Forward)",
            data=excel_fwd_bytes,
            file_name=nombre_fwd_excel,
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            key="btn_descarga_forward"
        )

        st.markdown("---")
        
        # === GENERADOR EXCEL PARA PROPHET ===
        st.subheader("⚙️ Generador de excel para Prophet")
        st.markdown("Genera una tabla con tasas forward requerida por Prophet.")
        
        col_p1, col_p2, col_p3 = st.columns(3)
        with col_p1:
            prophet_fecha = st.date_input("Fecha de evaluación", value=datetime.date(2024, 12, 31), key="prop_fecha")
        with col_p2:
            prophet_producto = st.selectbox("Tipo de producto", 
                ["Vida Tradicional", "Ahorro-Inversión", "SPP y SCTR", "Renta Particular"], 
                key="prop_prod"
            )
        with col_p3:
            prophet_mes_inicio = st.selectbox("Mes de Inicio", ["Mes 0", "Mes 1"], key="prop_mes")
            
        st.info(f"💡 Se generarán las tasas forward para **{prophet_producto}** a partir de las tasas spot (inc. VOLA).")
            
        if st.button("🚀 Generar Excel Prophet", type="primary"):
            excel_prophet_bytes = generar_excel_prophet(
                dict_tasas=dict_3_monedas_forward,
                fecha_evaluacion=prophet_fecha,
                tipo_producto=prophet_producto,
                mes_inicio=prophet_mes_inicio
            )
            st.download_button(
                label="📥 Descargar Excel VTD para Prophet",
                data=excel_prophet_bytes,
                file_name=f"Prophet_VTD_{prophet_producto.replace(' ', '_')}.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                key="btn_descarga_prophet"
            )
        
        st.markdown("---")

        cols_fwd = st.columns(len(dict_forward))
        for i, (moneda, df) in enumerate(dict_forward.items()):
            with cols_fwd[i]:
                st.markdown(f"**{moneda}**")
                df_show = df[['Año', 'Mes', 'Tasa anual + VOLA (%)', 'Tasa Forward (%)']]
                st.dataframe(
                    df_show.style.format({
                        'Tasa anual + VOLA (%)': '{:.4%}',
                        'Tasa Forward (%)': '{:.4%}'
                    }),
                    hide_index=True,
                    height=400
                )

# ==========================================
# PESTAÑA 3: COMPARATIVA DE PERIODOS
# ==========================================
with tab_comparativa:
    st.subheader("Comparar VTDA en múltiples periodos")

    moneda_comp = st.selectbox("Seleccione la moneda a evaluar:", list(archivos_fuente.keys()), key="moneda_comp")
    vola_default_comp = VOLA_VALORES.get(moneda_comp, 0.001502) * 100

    col1, col2, col3 = st.columns(3)
    periodos_config = []

    with col1:
        st.markdown("##### 📅 Periodo 1")
        p1_inicio = st.date_input("Inicio P1", value=datetime.date(2024, 1, 1), min_value=datetime.date(2015, 1, 1), max_value=datetime.date(2030, 12, 31), key="p1_ini")
        p1_fin = st.date_input("Fin P1", value=datetime.date(2024, 12, 31), min_value=datetime.date(2015, 1, 1), max_value=datetime.date(2030, 12, 31), key="p1_fin")
        p1_vola = st.number_input("VOLA P1 (%)", value=vola_default_comp, format="%.4f", key="vola_p1") / 100
        periodos_config.append(("Periodo 1", p1_inicio, p1_fin, p1_vola))

    with col2:
        st.markdown("##### 📅 Periodo 2")
        p2_act = st.checkbox("Habilitar Periodo 2", value=True)
        p2_inicio = st.date_input("Inicio P2", value=datetime.date(2025, 1, 1), min_value=datetime.date(2015, 1, 1), max_value=datetime.date(2030, 12, 31), key="p2_ini")
        p2_fin = st.date_input("Fin P2", value=datetime.date(2025, 12, 31), min_value=datetime.date(2015, 1, 1), max_value=datetime.date(2030, 12, 31), key="p2_fin")
        p2_vola = st.number_input("VOLA P2 (%)", value=vola_default_comp, format="%.4f", key="vola_p2") / 100
        if p2_act:
            periodos_config.append(("Periodo 2", p2_inicio, p2_fin, p2_vola))

    with col3:
        st.markdown("##### 📅 Periodo 3")
        p3_act = st.checkbox("Habilitar Periodo 3", value=False)
        p3_inicio = st.date_input("Inicio P3", value=datetime.date(2025, 6, 1), min_value=datetime.date(2015, 1, 1), max_value=datetime.date(2030, 12, 31), key="p3_ini")
        p3_fin = st.date_input("Fin P3", value=datetime.date(2025, 12, 31), min_value=datetime.date(2015, 1, 1), max_value=datetime.date(2030, 12, 31), key="p3_fin")
        p3_vola = st.number_input("VOLA P3 (%)", value=vola_default_comp, format="%.4f", key="vola_p3") / 100
        if p3_act:
            periodos_config.append(("Periodo 3", p3_inicio, p3_fin, p3_vola))

    st.markdown("---")

    if st.button("Generar Comparativa", type="primary"):
        st.session_state.comp_df_list = []
        st.session_state.comp_resultados = {}
        
        if moneda_comp in bases_de_datos and bases_de_datos[moneda_comp] is not None:
            for etiqueta, p_ini, p_fin, p_vola in periodos_config:
                df_p = calcular_curva_cupon_cero(bases_de_datos[moneda_comp], p_ini, p_fin)
                if df_p is not None:
                    df_p['Tasa anual + VOLA (%)'] = df_p['Tasa anual (%)'] + p_vola
                    nombre_periodo = f"{etiqueta} ({p_ini.strftime('%d/%m/%y')} - {p_fin.strftime('%d/%m/%y')})"
                    df_p['Periodo'] = nombre_periodo

                    st.session_state.comp_df_list.append(df_p)
                    st.session_state.comp_resultados[etiqueta] = df_p

            if not st.session_state.comp_df_list:
                st.warning("⚠️ No se encontraron datos suficientes para los periodos seleccionados.")
        else:
            st.error("Los datos de esta moneda no están disponibles o no se pudieron cargar.")

    if st.session_state.comp_df_list:
        df_comparativo = pd.concat(st.session_state.comp_df_list)
        df_comparativo['Tasa Graficar (%)'] = df_comparativo['Tasa anual + VOLA (%)'] * 100

        vista_eje_comp = st.radio(
            "Seleccione la escala del gráfico:",
            ["Mensual", "Anual"],
            horizontal=True,
            key="vista_eje_comp"
        )

        if vista_eje_comp == "Mensual":
            df_comp_final = df_comparativo[df_comparativo['Mes'] <= 240]
            col_eje_comp = "Mes"
            etiqueta_x_comp = "Plazo (Meses)"
        else:
            df_comp_final = df_comparativo[(df_comparativo['Mes'] % 12 == 0) & (df_comparativo['Año'] <= 20)]
            col_eje_comp = "Año"
            etiqueta_x_comp = "Plazo (Años)"

        fig_comp = px.line(
            df_comp_final,
            x=col_eje_comp,
            y="Tasa Graficar (%)",
            color="Periodo",
            title=f"Comparativa Curva Cupón Cero - {moneda_comp} (Tasa Anual + VOLA)",
            labels={"Tasa Graficar (%)": "Tasa Anual + VOLA (%)", col_eje_comp: etiqueta_x_comp},
            template="plotly_white",
            color_discrete_sequence=colores_graficos,
            markers=True if col_eje_comp == "Año" else False
        )
        fig_comp.update_layout(yaxis_tickformat='.2f', legend=dict(orientation="h", yanchor="bottom", y=1.02, xanchor="right", x=1), margin=dict(l=10, r=10, t=50, b=10))

        if col_eje_comp == "Año":
            fig_comp.update_xaxes(dtick=1, range=[0, 20])
        else:
            fig_comp.update_xaxes(dtick=12, range=[0, 240])

        st.plotly_chart(fig_comp, use_container_width=True)

        cols_tables = st.columns(len(st.session_state.comp_resultados))
        for idx, (etiqueta, df_res) in enumerate(st.session_state.comp_resultados.items()):
            with cols_tables[idx]:
                nombre_periodo = df_res['Periodo'].iloc[0]
                st.markdown(f"**{nombre_periodo}**")
                df_show = df_res[['Año', 'Mes', 'Tasa anual + VOLA (%)']]
                st.dataframe(df_show.style.format({'Tasa anual + VOLA (%)': '{:.4%}'}), hide_index=True, height=250)

        st.markdown("---")
        nombre_moneda_corta = moneda_comp.split()[0]
        nombre_excel_comp = f"Comparativa_VTD_{nombre_moneda_corta}.xlsx"

        buffer_comp = io.BytesIO()
        with pd.ExcelWriter(buffer_comp, engine='openpyxl') as writer:
            df_resumen_data = {
                'Parámetro': ['Moneda Evaluada'] + [f'Periodo {i+1}' for i in range(len(st.session_state.comp_resultados))],
                'Valor': [moneda_comp] + [df['Periodo'].iloc[0] for df in st.session_state.comp_resultados.values()]
            }
            df_resumen_comp = pd.DataFrame(df_resumen_data)
            df_resumen_comp.to_excel(writer, sheet_name='Resumen_Filtros', index=False)

            for etiqueta, df_res in st.session_state.comp_resultados.items():
                df_export = df_res.copy()
                if 'Periodo' in df_export.columns:
                    df_export.drop(columns=['Periodo'], inplace=True)
                df_export.to_excel(writer, sheet_name=etiqueta.replace(" ", "_"), index=False)

        st.download_button(
            label="📥 Descargar Comparativa en formato Excel",
            data=buffer_comp.getvalue(),
            file_name=nombre_excel_comp,
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            key="btn_descarga_comparativa"
        )

# ==========================================
# PESTAÑA 4: HISTÓRICO VOLA
# ==========================================
with tab_vola_hist:
    st.header("Tasas de Ajuste por Volatilidad (VOLA)")
    st.markdown("""
    Esta sección consolida las tasas VOLA históricas empleadas para distintos requerimientos regulatorios. 
    Permite visualizar las tasas para reservas matemáticas y para los indicadores de Gestión de Activos y Pasivos (GAP).
    """)
    
    archivo_vola = "Vola_Historico.xlsx"
    
    df_vola_reservas = None
    df_vola_spp = None
    df_vola_vlp = None

    try:
        if os.path.exists(archivo_vola):
            df_vola_reservas = pd.read_excel(archivo_vola, sheet_name="VOLA (no SPP-SCTR)", skiprows=2)
            df_vola_spp = pd.read_excel(archivo_vola, sheet_name="VOLA SPP-SCTR", skiprows=2)
            df_vola_vlp = pd.read_excel(archivo_vola, sheet_name="VOLA VLP", skiprows=2)
            
            df_vola_reservas = df_vola_reservas.loc[:, ~df_vola_reservas.columns.str.contains('^Unnamed')]
            df_vola_spp = df_vola_spp.loc[:, ~df_vola_spp.columns.str.contains('^Unnamed')]
            df_vola_vlp = df_vola_vlp.loc[:, ~df_vola_vlp.columns.str.contains('^Unnamed')]
            
            meses_es = {1:'Ene', 2:'Feb', 3:'Mar', 4:'Abr', 5:'May', 6:'Jun', 
                        7:'Jul', 8:'Ago', 9:'Set', 10:'Oct', 11:'Nov', 12:'Dic'}
            
            def formato_mes_ano(x):
                if pd.isnull(x): return x
                return f"{meses_es[x.month]}-{x.strftime('%y')}"

            if 'Fecha de reporte' in df_vola_spp.columns:
                df_vola_spp['Fecha de reporte'] = pd.to_datetime(df_vola_spp['Fecha de reporte']).apply(formato_mes_ano)
                
            if 'Fecha de reporte' in df_vola_vlp.columns:
                df_vola_vlp['Fecha de reporte'] = pd.to_datetime(df_vola_vlp['Fecha de reporte']).apply(formato_mes_ano)

            st.success(f"✅ Datos históricos del VOLA cargados exitosamente.")
        else:
            st.info(f"ℹ️ El archivo '{archivo_vola}' no se encontró en la ruta raíz. Cargando la base de datos de respaldo predeterminada.")
            
            df_vola_reservas = pd.DataFrame({
                "Año de reporte": [2020, 2021, 2022, 2023, 2024, 2025, 2026],
                "Soles": [0.001996, 0.002245, 0.002408, 0.001856, 0.001868, 0.001040, 0.001308],
                "Soles VAC": [0.002424, 0.002600, 0.002130, 0.002386, 0.000522, 0.000636, 0.000003],
                "Dólares": [0.002274, 0.002389, 0.003134, 0.003553, 0.000951, 0.001083, 0.000862]
            })
            
            df_vola_spp = pd.DataFrame({
                "Fecha de reporte": ["Mar-26", "Jun-26"],
                "Soles": [0.000907, 0.001502],
                "Soles VAC": [0.001449, 0.002471],
                "Dólares": [0.001389, 0.002175]
            })
            
            df_vola_vlp = pd.DataFrame({
                "Fecha de reporte": ["Mar-26", "Jun-26"],
                "Soles": [0.001430, 0.001853],
                "Soles VAC": [0.001028, 0.000003],
                "Dólares": [0.000889, 0.001558]
            })
    except Exception as e:
        st.error(f"Ocurrió un error al cargar los datos históricos de VOLA: {e}")

    if df_vola_reservas is not None:
        st.markdown("---")
        
        st.subheader("📊 Reservas Matemáticas (Seguros de Vida no SPP/SCTR)")
        
        col_chart, col_table = st.columns([6, 4])

        with col_chart:
            # Opción para elegir el periodo a graficar
            filtro_res = st.radio(
                "Periodo a visualizar (Reservas):",
                ["Últimos 10 años", "Histórico completo"],
                horizontal=True,
                key="rad_res"
            )
            
            # Filtrar el dataframe según la selección
            df_res_plot = df_vola_reservas.tail(10) if filtro_res == "Últimos 10 años" else df_vola_reservas.copy()
            
            df_res_melt = df_res_plot.melt(id_vars="Año de reporte", var_name="Moneda", value_name="Tasa")
            df_res_melt["Tasa (%)"] = df_res_melt["Tasa"] * 100
            
            fig_res = px.line(
                df_res_melt,
                x="Año de reporte",
                y="Tasa (%)",
                color="Moneda",
                markers=True,
                title="Evolución Histórica VOLA - Reservas Matemáticas",
                color_discrete_sequence=["#06369D", "#ED8B00", "#1FC3B3"]
            )
            fig_res.update_layout(
                yaxis_tickformat='.3f',
                legend=dict(orientation="h", yanchor="bottom", y=1.02, xanchor="right", x=1),
                margin=dict(l=10, r=10, t=40, b=10)
            )
            fig_res.update_xaxes(dtick=1)
            st.plotly_chart(fig_res, use_container_width=True)


        with col_table:
            st.markdown("<br>", unsafe_allow_html=True)
            st.dataframe(
                df_vola_reservas.style.format({
                    "Soles": "{:.4%}", "Soles VAC": "{:.4%}", "Dólares": "{:.4%}"
                }),
                hide_index=True,
                use_container_width=True
            )

        st.markdown("---")

        st.subheader("📈 Gestión de Activos y Pasivos")
        
        # Selector global para los gráficos de GAP (SPP/SCTR y VLP)
        filtro_gap = st.radio(
            "Periodo a visualizar (GAP):",
            ["Últimos 10 trimestres", "Histórico completo"],
            horizontal=True,
            key="rad_gap"
        )
        ver_todo_gap = (filtro_gap == "Histórico completo")
        
        def mostrar_kpis(df, titulo, mostrar_todo=False):
            st.markdown(f"**{titulo}**")
            if len(df) >= 2:
                ult = df.iloc[-1]
                ant = df.iloc[-2]
                
                ayuda_pbs = "VOLA del último trimestre registrado. Y debajo está la variación respecto al trimestre previo (100 pbs = 1%)."
                
                kpi1, kpi2, kpi3 = st.columns(3)
                with kpi1:
                    st.metric("Soles", f"{ult['Soles']:.4%}", f"{(ult['Soles'] - ant['Soles']) * 10000:.1f} pbs", help=ayuda_pbs)
                with kpi2:
                    st.metric("Soles VAC", f"{ult['Soles VAC']:.4%}", f"{(ult['Soles VAC'] - ant['Soles VAC']) * 10000:.1f} pbs", help=ayuda_pbs)
                with kpi3:
                    st.metric("Dólares", f"{ult['Dólares']:.4%}", f"{(ult['Dólares'] - ant['Dólares']) * 10000:.1f} pbs", help=ayuda_pbs)
            
            # --- Gráfico de tendencia filtrado ---
            if not df.empty and 'Fecha de reporte' in df.columns:
                # Aplicar el filtro de los últimos 10 registros (trimestres)
                df_plot = df.copy() if mostrar_todo else df.tail(10)
                
                df_melt = df_plot.melt(id_vars="Fecha de reporte", var_name="Moneda", value_name="Tasa")
                df_melt["Tasa (%)"] = df_melt["Tasa"] * 100
                
                fig = px.line(
                    df_melt,
                    x="Fecha de reporte",
                    y="Tasa (%)",
                    color="Moneda",
                    markers=True,
                    color_discrete_sequence=["#06369D", "#ED8B00", "#1FC3B3"]
                )
                fig.update_layout(
                    yaxis_tickformat='.3f',
                    legend=dict(orientation="h", yanchor="bottom", y=1.02, xanchor="right", x=1),
                    margin=dict(l=10, r=10, t=10, b=10),
                    xaxis_title=None,
                    yaxis_title="Tasa (%)",
                    height=250
                )
                st.plotly_chart(fig, use_container_width=True)
            
            st.dataframe(
                df.style.format({"Soles": "{:.4%}", "Soles VAC": "{:.4%}", "Dólares": "{:.4%}"}),
                hide_index=True, use_container_width=True
            )
 
        col_spp, col_vlp = st.columns(2)
        
        with col_spp:
            mostrar_kpis(df_vola_spp, "VOLA para SPP y SCTR", mostrar_todo=ver_todo_gap)
            
        with col_vlp:
            mostrar_kpis(df_vola_vlp, "VOLA para Vida Largo Plazo (no SPP/SCTR)", mostrar_todo=ver_todo_gap)
            st.caption("*(i) Este VOLA se emplea para los GHO regulatorios de rentas particulares, y seguros de vida largo plazo, tradicional y con componente de ahorro y/o inversión.*")


# ==========================================
# PESTAÑA 5: CÁLCULO DE VPN (ANÁLISIS DE SENSIBILIDAD)
# ==========================================
with tab_vpn:
    st.header("Valor Presente Neto (VPN)")
    st.markdown("Calcula el VPN de un flujo de caja comparando escenarios generados en el dashboard o usando una curva manual.")
 
    col_flujo, col_tasa = st.columns([1, 1])
    
    with col_flujo:
        st.subheader("1. Cargar flujo mensual")
        archivo_flujo = st.file_uploader("Sube el Excel con el flujo (columnas: Periodo, Flujo)", type=["xlsx", "xls"], key="up_flujo_vpn")
    
    with col_tasa:
        st.subheader("2. Origen del vector de tasas")
        origen_vtd = st.radio(
            "Selecciona cómo definir la(s) curva(s):",
            ["Generar escenarios múltiples", "Cargar VTD manual"]
        )
 
    # Variables para almacenar configuración
    escenarios_vpn = []
    df_vtd_manual = None
 
    st.markdown("---")
 
    # --- Lógica de Interfaz según selección ---
    if origen_vtd == "Generar escenarios múltiples":
        st.subheader("3. Configurar Escenarios de Tasas")
        moneda_vpn = st.selectbox("Seleccione la moneda a evaluar:", list(archivos_fuente.keys()), key="moneda_vpn")
        vola_default_vpn = VOLA_VALORES.get(moneda_vpn, 0.001502) * 100
 
        col_p1, col_p2, col_p3 = st.columns(3)
 
        with col_p1:
            st.markdown("##### 📅 Escenario 1")
            e1_inicio = st.date_input("Inicio E1", value=datetime.date(2024, 1, 1), key="e1_ini")
            e1_fin = st.date_input("Fin E1", value=datetime.date(2024, 12, 31), key="e1_fin")
            e1_vola = st.number_input("VOLA E1 (%)", value=vola_default_vpn, format="%.4f", key="vola_e1") / 100
            escenarios_vpn.append(("Escenario 1", e1_inicio, e1_fin, e1_vola))
 
        with col_p2:
            st.markdown("##### 📅 Escenario 2")
            e2_act = st.checkbox("Habilitar Escenario 2", value=True, key="e2_act")
            e2_inicio = st.date_input("Inicio E2", value=datetime.date(2025, 1, 1), key="e2_ini")
            e2_fin = st.date_input("Fin E2", value=datetime.date(2025, 12, 31), key="e2_fin")
            e2_vola = st.number_input("VOLA E2 (%)", value=vola_default_vpn, format="%.4f", key="vola_e2") / 100
            if e2_act:
                escenarios_vpn.append(("Escenario 2", e2_inicio, e2_fin, e2_vola))
 
        with col_p3:
            st.markdown("##### 📅 Escenario 3")
            e3_act = st.checkbox("Habilitar Escenario 3", value=False, key="e3_act")
            e3_inicio = st.date_input("Inicio E3", value=datetime.date(2025, 6, 1), key="e3_ini")
            e3_fin = st.date_input("Fin E3", value=datetime.date(2025, 12, 31), key="e3_fin")
            e3_vola = st.number_input("VOLA E3 (%)", value=vola_default_vpn, format="%.4f", key="vola_e3") / 100
            if e3_act:
                escenarios_vpn.append(("Escenario 3", e3_inicio, e3_fin, e3_vola))
 
    else:
        st.subheader("3. Cargar archivo VTD")
        archivo_vtd = st.file_uploader("Sube el Excel con VTD (columnas: Periodo, VTDA)", type=["xlsx", "xls"], key="up_vtd")
        if archivo_vtd:
            try:
                df_vtd_manual = pd.read_excel(archivo_vtd)
                # Normalizar columnas por si tienen otros nombres
                if 'Periodo' not in df_vtd_manual.columns or 'VTDA' not in df_vtd_manual.columns:
                    df_vtd_manual.columns = ['Periodo', 'VTDA'] + list(df_vtd_manual.columns[2:])
                
                # Convertir porcentajes si se detectan mayores a 1
                if df_vtd_manual['VTDA'].mean() > 1:
                    df_vtd_manual['VTDA'] = df_vtd_manual['VTDA'] / 100
            except Exception as e:
                st.error(f"Error al leer archivo VTD: {e}")
 
    # --- Lógica de Cálculo ---
    if archivo_flujo is not None:
        st.markdown("---")
        if st.button("🚀 Calcular VPN", type="primary", use_container_width=True):
            try:
                df_flujo = pd.read_excel(archivo_flujo)
                # Normalizar columnas de flujo
                if 'Periodo' not in df_flujo.columns or 'Flujo' not in df_flujo.columns:
                    df_flujo.columns = ['Periodo', 'Flujo'] + list(df_flujo.columns[2:])
 
                resultados_vpn = []
                df_detalles = df_flujo[['Periodo', 'Flujo']].copy()
                extrapolado_flag = False
 
                if origen_vtd == "Generar escenarios múltiples":
                    if moneda_vpn in bases_de_datos and bases_de_datos[moneda_vpn] is not None:
                        for etiqueta, p_ini, p_fin, p_vola in escenarios_vpn:
                            df_curva = calcular_curva_cupon_cero(bases_de_datos[moneda_vpn], p_ini, p_fin)
                            
                            if df_curva is not None:
                                df_curva['VTDA'] = df_curva['Tasa anual (%)'] + p_vola
                                
                                # Merge de flujo con la curva generada
                                df_calc = pd.merge(df_flujo[['Periodo', 'Flujo']], df_curva[['Mes', 'VTDA']], left_on='Periodo', right_on='Mes', how='left')
                                
                                if df_calc['VTDA'].isna().any():
                                    df_calc['VTDA'] = df_calc['VTDA'].ffill().bfill()
                                    extrapolado_flag = True
                                    
                                # Descuento y VPN Spot Anualizado
                                df_calc['Factor Descuento'] = 1 / ((1 + df_calc['VTDA']) ** (df_calc['Periodo'] / 12))
                                df_calc['Flujo Descontado'] = df_calc['Flujo'] * df_calc['Factor Descuento']
                                vpn_total = df_calc['Flujo Descontado'].sum()
                                
                                df_detalles[f"VTDA {etiqueta}"] = df_calc['VTDA']
                                df_detalles[f"Flujo Desc. {etiqueta}"] = df_calc['Flujo Descontado']
                                
                                resultados_vpn.append({
                                    "Escenario": f"{etiqueta}\n({p_ini.strftime('%m/%y')} - {p_fin.strftime('%m/%y')})",
                                    "VPN Total": vpn_total
                                })
                            else:
                                st.warning(f"No se pudo generar la curva para el {etiqueta}.")
                else:
                    if df_vtd_manual is not None:
                        df_calc = pd.merge(df_flujo[['Periodo', 'Flujo']], df_vtd_manual[['Periodo', 'VTDA']], on='Periodo', how='left')
                        
                        if df_calc['VTDA'].isna().any():
                            df_calc['VTDA'] = df_calc['VTDA'].ffill().bfill()
                            extrapolado_flag = True
 
                        df_calc['Factor Descuento'] = 1 / ((1 + df_calc['VTDA']) ** (df_calc['Periodo'] / 12))
                        df_calc['Flujo Descontado'] = df_calc['Flujo'] * df_calc['Factor Descuento']
                        vpn_total = df_calc['Flujo Descontado'].sum()
                        
                        df_detalles["VTDA Manual"] = df_calc['VTDA']
                        df_detalles["Flujo Desc. Manual"] = df_calc['Flujo Descontado']
                        
                        resultados_vpn.append({
                            "Escenario": "VTD Manual Cargada",
                            "VPN Total": vpn_total
                        })
                    else:
                        st.warning("⚠️ Debes cargar un archivo Excel válido con la VTD.")
 
                if extrapolado_flag:
                    st.toast("💡 El flujo excede los periodos de la VTD. Se ha extrapolado de forma plana la última tasa.")


                # --- Renderizado de Resultados ---
                if resultados_vpn:
                    df_res_vpn = pd.DataFrame(resultados_vpn)
                    
                    st.subheader("📊 Resultados de VPN")
                    
                    # 1. Cálculo dinámico del eje Y (para no empezar desde 0 y resaltar diferencias)
                    min_vpn = df_res_vpn["VPN Total"].min()
                    max_vpn = df_res_vpn["VPN Total"].max()
                    dif_vpn = max_vpn - min_vpn
                    
                    # Margen inferior y superior (dejamos más espacio arriba para que el texto no se corte)
                    margen = dif_vpn * 0.3 if dif_vpn > 0 else abs(min_vpn) * 0.1
                    rango_y = [min_vpn - margen, max_vpn + (margen * 1.5)]
                    
                    # 2. Creación del gráfico
                    fig_bar = px.bar(
                        df_res_vpn,
                        x="Escenario",
                        y="VPN Total",
                        text="VPN Total",
                        color="Escenario",
                        color_discrete_sequence=colores_graficos
                    )
                    
                    # 3. Ajuste de barras (más delgadas) y posición del texto
                    fig_bar.update_traces(
                        texttemplate='%{text:,.2f}',
                        textposition='outside',
                        textfont_size=12,
                        width=0.4  # Hace las barras más angostas
                    )
                    
                    # 4. Limpieza visual general
                    fig_bar.update_layout(
                        height=350,               # Gráfico más pequeño en altura
                        showlegend=False,
                        plot_bgcolor="rgba(0,0,0,0)",  # Fondo transparente
                        margin=dict(t=40, b=10, l=10, r=10),
                        yaxis=dict(
                            title="",             # Quitamos título del eje Y
                            tickformat=',.2f',
                            showgrid=False,       # Quitamos líneas guía horizontales
                            zeroline=False,       # Quitamos la línea de origen cero
                            range=rango_y         # Aplicamos el rango dinámico
                        ),
                        xaxis=dict(
                            title="",             # Quitamos título del eje X
                            showgrid=False,       # Quitamos líneas guía verticales
                            zeroline=False
                        )
                    )
                    
                    st.plotly_chart(fig_bar, use_container_width=True)
                    
                    cols_kpi = st.columns(len(resultados_vpn))
                    for i, res in enumerate(resultados_vpn):
                        with cols_kpi[i]:
                            st.markdown(
                                f"""
                                <div style="background-color: #EAF0F8; padding: 15px; border-radius: 5px; border-left: 6px solid {colores_graficos[i % len(colores_graficos)]}; height: 100%;">
                                    <h5 style="color: {colores_graficos[i % len(colores_graficos)]}; margin: 0; font-size: 14px;">{res['Escenario'].replace(chr(10), ' ')}</h5>
                                    <h3 style="margin: 5px 0 0 0; font-size: 22px;">{res['VPN Total']:,.2f}</h3>
                                </div>
                                """,
                                unsafe_allow_html=True
                            )
 
                    st.markdown("---")
                    st.subheader("📋 Detalle de Flujos y Descuentos")
                    
                    # Formato dinámico de las columnas
                    format_dict = {'Flujo': '{:,.2f}'}
                    for col in df_detalles.columns:
                        if "VTDA" in col:
                            format_dict[col] = '{:.4%}'
                        elif "Flujo Desc" in col:
                            format_dict[col] = '{:,.2f}'
 
                    st.dataframe(df_detalles.style.format(format_dict), use_container_width=True)
                    
                    # Botón de descarga unificada
                    buffer_vpn = io.BytesIO()
                    with pd.ExcelWriter(buffer_vpn, engine='openpyxl') as writer:
                        df_res_vpn.to_excel(writer, index=False, sheet_name='Resumen_VPN')
                        df_detalles.to_excel(writer, index=False, sheet_name='Detalle_Flujos')
                    
                    st.download_button(
                        label="📥 Descargar Detalle del VPN en Excel",
                        data=buffer_vpn.getvalue(),
                        file_name="Analisis_Sensibilidad_VPN.xlsx",
                        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                    )
 
            except Exception as e:
                st.error(f"Error en el cálculo: Verifica la estructura de tus archivos. Detalle: {e}")